import csv
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db import transaction, models
from django.db.models import Q
from django.contrib.auth import get_user_model
from django.core.paginator import Paginator
from openpyxl import load_workbook
from academics.models import Department, Program, Semester, AcademicSession
from .models import Student, StudentAcademicHistory, Enrollment
from academics.views import admin_required
from audit.models import AuditLog

User = get_user_model()

@login_required
def student_list(request):
    inst = request.user.institution
    students = Student.objects.filter(institution=inst).select_related('user', 'department', 'program', 'semester')
    
    # Filters
    dept_id = request.GET.get('department')
    prog_id = request.GET.get('program')
    sem_id = request.GET.get('semester')
    q = request.GET.get('q', '')
    
    if dept_id:
        students = students.filter(department_id=dept_id)
    if prog_id:
        students = students.filter(program_id=prog_id)
    if sem_id:
        students = students.filter(semester_id=sem_id)
    if q:
        students = students.filter(
            Q(user__first_name__icontains=q) | 
            Q(user__last_name__icontains=q) |
            Q(student_id__icontains=q) |
            Q(registration_number__icontains=q)
        )
        
    paginator = Paginator(students, 10)  # 10 per page
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)
    
    context = {
        'page_obj': page_obj,
        'departments': Department.objects.filter(institution=inst),
        'programs': Program.objects.filter(institution=inst),
        'semesters': Semester.objects.filter(institution=inst),
        'dept_filter': int(dept_id) if dept_id else '',
        'prog_filter': int(prog_id) if prog_id else '',
        'sem_filter': int(sem_id) if sem_id else '',
        'q': q,
        'breadcrumbs': [{'name': 'Students', 'url': '#'}]
    }
    return render(request, 'students/student_list.html', context)

@login_required
@admin_required
def student_add(request):
    inst = request.user.institution
    depts = Department.objects.filter(institution=inst)
    programs = Program.objects.filter(institution=inst)
    semesters = Semester.objects.filter(institution=inst)
    
    if request.method == 'POST':
        username = request.POST.get('username')
        email = request.POST.get('email')
        password = request.POST.get('password')
        first_name = request.POST.get('first_name')
        last_name = request.POST.get('last_name')
        student_id = request.POST.get('student_id')
        reg_num = request.POST.get('registration_number')
        dob = request.POST.get('dob')
        gender = request.POST.get('gender')
        father = request.POST.get('father_name')
        mother = request.POST.get('mother_name')
        parent_contact = request.POST.get('parent_contact')
        
        dept_id = request.POST.get('department')
        prog_id = request.POST.get('program')
        sem_id = request.POST.get('semester')
        
        if username and email and password and student_id and dept_id and prog_id and sem_id:
            dept = get_object_or_404(Department, id=dept_id, institution=inst)
            prog = get_object_or_404(Program, id=prog_id, institution=inst)
            sem = get_object_or_404(Semester, id=sem_id, program=prog)
            
            try:
                with transaction.atomic():
                    # Create User
                    user = User.objects.create_user(
                        username=username,
                        email=email,
                        password=password,
                        first_name=first_name,
                        last_name=last_name,
                        role=User.Role.STUDENT,
                        institution=inst
                    )
                    
                    # Create Student
                    student = Student.objects.create(
                        institution=inst,
                        user=user,
                        student_id=student_id,
                        registration_number=reg_num,
                        dob=dob if dob else None,
                        gender=gender,
                        father_name=father,
                        mother_name=mother,
                        parent_contact=parent_contact,
                        department=dept,
                        program=prog,
                        semester=sem,
                        academic_session=inst.settings.current_session
                    )
                    
                    # Create Academic History
                    StudentAcademicHistory.objects.create(student=student)
                    
                    # Create initial Enrollment
                    if inst.settings.current_session:
                        Enrollment.objects.create(
                            student=student,
                            academic_session=inst.settings.current_session,
                            semester=sem
                        )
                        
                    AuditLog.objects.create(
                        institution=inst,
                        user=request.user,
                        user_role=request.user.get_role_display(),
                        action=f"Created student profile: {student_id}.",
                        module="STUDENTS",
                        ip_address=request.META.get('REMOTE_ADDR')
                    )
                    
                messages.success(request, f"Student '{first_name} {last_name}' registered successfully.")
                if 'add_another' in request.POST:
                    return redirect('student_add')
                return redirect('student_list')
            except Exception as e:
                messages.error(request, f"Failed to register student: {e}")
        else:
            messages.error(request, "All required fields must be completed.")
            
    return render(request, 'students/student_form.html', {
        'title': 'Add Student',
        'departments': depts,
        'programs': programs,
        'semesters': semesters,
        'breadcrumbs': [{'name': 'Students', 'url': '/students/'}, {'name': 'Add Student', 'url': '#'}]
    })

@login_required
@admin_required
def student_edit(request, id):
    inst = request.user.institution
    student = get_object_or_404(Student, id=id, institution=inst)
    depts = Department.objects.filter(institution=inst)
    programs = Program.objects.filter(institution=inst)
    semesters = Semester.objects.filter(institution=inst)
    
    if request.method == 'POST':
        student.user.first_name = request.POST.get('first_name')
        student.user.last_name = request.POST.get('last_name')
        student.user.email = request.POST.get('email')
        student.user.save()
        
        student.registration_number = request.POST.get('registration_number')
        dob = request.POST.get('dob')
        student.dob = dob if dob else None
        student.gender = request.POST.get('gender')
        student.father_name = request.POST.get('father_name')
        student.mother_name = request.POST.get('mother_name')
        student.parent_contact = request.POST.get('parent_contact')
        
        dept_id = request.POST.get('department')
        prog_id = request.POST.get('program')
        sem_id = request.POST.get('semester')
        
        student.department = get_object_or_404(Department, id=dept_id, institution=inst)
        student.program = get_object_or_404(Program, id=prog_id, institution=inst)
        student.semester = get_object_or_404(Semester, id=sem_id, program=student.program)
        student.save()
        
        AuditLog.objects.create(
            institution=inst,
            user=request.user,
            user_role=request.user.get_role_display(),
            action=f"Edited student profile: {student.student_id}.",
            module="STUDENTS",
            ip_address=request.META.get('REMOTE_ADDR')
        )
        
        messages.success(request, "Student profile updated successfully.")
        return redirect('student_list')
        
    return render(request, 'students/student_form.html', {
        'title': 'Edit Student',
        'student': student,
        'departments': depts,
        'programs': programs,
        'semesters': semesters,
        'breadcrumbs': [{'name': 'Students', 'url': '/students/'}, {'name': 'Edit Student', 'url': '#'}]
    })

@login_required
def student_profile(request, id):
    inst = request.user.institution
    student = get_object_or_404(Student, id=id, institution=inst)
    
    # Security: A student user can only see their own profile
    if request.user.role == User.Role.STUDENT and request.user.student.id != student.id:
        messages.error(request, "Access denied.")
        return redirect('dashboard')
        
    return render(request, 'students/student_profile.html', {
        'student': student,
        'breadcrumbs': [{'name': 'Students', 'url': '/students/'}, {'name': 'Profile', 'url': '#'}]
    })

@login_required
@admin_required
def student_promotion(request):
    inst = request.user.institution
    programs = Program.objects.filter(institution=inst)
    semesters = Semester.objects.filter(institution=inst)
    
    if request.method == 'POST':
        prog_id = request.POST.get('program')
        current_sem_num = request.POST.get('current_semester')
        
        if prog_id and current_sem_num:
            prog = get_object_or_404(Program, id=prog_id, institution=inst)
            current_sem = get_object_or_404(Semester, program=prog, number=int(current_sem_num))
            
            # Find next semester
            next_sem_num = current_sem.number + 1
            next_sem = Semester.objects.filter(program=prog, number=next_sem_num).first()
            
            if not next_sem:
                messages.error(request, f"There is no next semester after Semester {current_sem_num} for this program.")
                return redirect('student_promotion')
                
            # Promote all students in this program and semester
            students_to_promote = Student.objects.filter(
                institution=inst, 
                program=prog, 
                semester=current_sem,
                status=Student.Status.ACTIVE
            )
            
            count = students_to_promote.count()
            if count == 0:
                messages.warning(request, "No active students found in the selected semester.")
                return redirect('student_promotion')
                
            with transaction.atomic():
                for student in students_to_promote:
                    student.semester = next_sem
                    student.save()
                    # Add enrollment tracking
                    Enrollment.objects.get_or_create(
                        student=student,
                        academic_session=inst.settings.current_session,
                        semester=next_sem
                    )
                
                AuditLog.objects.create(
                    institution=inst,
                    user=request.user,
                    user_role=request.user.get_role_display(),
                    action=f"Promoted {count} students from Semester {current_sem_num} to Semester {next_sem_num} in program {prog.code}.",
                    module="STUDENTS",
                    ip_address=request.META.get('REMOTE_ADDR')
                )
                
            messages.success(request, f"Successfully promoted {count} students to Semester {next_sem_num}!")
            return redirect('student_list')
            
    return render(request, 'students/student_promotion.html', {
        'programs': programs,
        'semesters': semesters,
        'breadcrumbs': [{'name': 'Students', 'url': '/students/'}, {'name': 'Bulk Promotion', 'url': '#'}]
    })

@login_required
@admin_required
def student_import(request):
    inst = request.user.institution
    stage = 'upload'
    errors = []
    preview_records = []
    
    if request.method == 'POST':
        action = request.POST.get('action')
        
        if action == 'upload_file':
            excel_file = request.FILES.get('file')
            if not excel_file:
                messages.error(request, "Please choose a file to upload.")
                return redirect('student_import')
                
            filename = excel_file.name
            if not (filename.endswith('.xlsx') or filename.endswith('.csv')):
                messages.error(request, "Unsupported file format. Please upload .xlsx or .csv files.")
                return redirect('student_import')
                
            try:
                # Store file data in session temporarily for multi-stage confirm
                data_list = []
                if filename.endswith('.xlsx'):
                    wb = load_workbook(excel_file, read_only=True)
                    sheet = wb.active
                    
                    # Read rows
                    headers = None
                    for row in sheet.iter_rows(values_only=True):
                        if not headers:
                            headers = [str(h).strip().lower() for h in row if h is not None]
                            continue
                        if any(row):  # If row is not completely empty
                            data_list.append(list(row))
                else:
                    # CSV processing
                    csv_data = csv.reader(excel_file.read().decode('utf-8').splitlines())
                    headers = [h.strip().lower() for h in next(csv_data)]
                    for row in csv_data:
                        if any(row):
                            data_list.append(row)
                
                # Check for required headers: username, email, password, student_id, first_name, last_name, dept_code, program_code, semester
                required_cols = ['username', 'email', 'password', 'student_id', 'first_name', 'last_name', 'dept_code', 'program_code', 'semester']
                missing_cols = [col for col in required_cols if col not in headers]
                if missing_cols:
                    messages.error(request, f"Missing required columns: {', '.join(missing_cols)}")
                    return redirect('student_import')
                
                # Map headers to indices
                col_map = {col: headers.index(col) for col in required_cols}
                
                # Validation loop
                validated_records = []
                for idx, row in enumerate(data_list, start=2):
                    row_errors = []
                    # Extracted variables
                    u_name = str(row[col_map['username']]).strip()
                    email = str(row[col_map['email']]).strip()
                    pw = str(row[col_map['password']]).strip()
                    s_id = str(row[col_map['student_id']]).strip()
                    f_name = str(row[col_map['first_name']]).strip()
                    l_name = str(row[col_map['last_name']]).strip()
                    d_code = str(row[col_map['dept_code']]).strip().upper()
                    p_code = str(row[col_map['program_code']]).strip().upper()
                    sem_num = int(row[col_map['semester']])
                    
                    # DB checks
                    if User.objects.filter(username=u_name).exists():
                        row_errors.append(f"Username '{u_name}' is already taken.")
                    if User.objects.filter(email=email).exists():
                        row_errors.append(f"Email '{email}' is already registered.")
                    if Student.objects.filter(institution=inst, student_id=s_id).exists():
                        row_errors.append(f"Roll Number '{s_id}' already exists in this institution.")
                        
                    dept = Department.objects.filter(institution=inst, code=d_code).first()
                    if not dept:
                        row_errors.append(f"Department code '{d_code}' not found.")
                        
                    prog = Program.objects.filter(institution=inst, code=p_code).first()
                    if not prog:
                        row_errors.append(f"Program code '{p_code}' not found.")
                        
                    sem = None
                    if prog:
                        sem = Semester.objects.filter(program=prog, number=sem_num).first()
                        if not sem:
                            row_errors.append(f"Semester {sem_num} does not exist for program '{p_code}'.")
                    
                    if row_errors:
                        errors.append({'row': idx, 'msgs': row_errors})
                    else:
                        validated_records.append({
                            'username': u_name, 'email': email, 'password': pw,
                            'student_id': s_id, 'first_name': f_name, 'last_name': l_name,
                            'dept_id': dept.id, 'prog_id': prog.id, 'sem_id': sem.id
                        })
                
                if errors:
                    stage = 'errors'
                else:
                    # Save validated records to session for Step 3 Preview/Confirm
                    request.session['validated_import_students'] = validated_records
                    preview_records = validated_records
                    stage = 'preview'
                    
            except Exception as e:
                messages.error(request, f"Error processing file: {e}")
                return redirect('student_import')
                
        elif action == 'confirm_import':
            records = request.session.get('validated_import_students')
            if not records:
                messages.error(request, "No validated records found to import.")
                return redirect('student_import')
                
            try:
                session_active = inst.settings.current_session
                if not session_active:
                    messages.error(request, "No active academic session configured for this institution. Configure one first.")
                    return redirect('student_import')
                    
                with transaction.atomic():
                    for rec in records:
                        user = User.objects.create_user(
                            username=rec['username'],
                            email=rec['email'],
                            password=rec['password'],
                            first_name=rec['first_name'],
                            last_name=rec['last_name'],
                            role=User.Role.STUDENT,
                            institution=inst
                        )
                        
                        student = Student.objects.create(
                            institution=inst,
                            user=user,
                            student_id=rec['student_id'],
                            department_id=rec['dept_id'],
                            program_id=rec['prog_id'],
                            semester_id=rec['sem_id'],
                            academic_session=session_active
                        )
                        
                        StudentAcademicHistory.objects.create(student=student)
                        
                        Enrollment.objects.create(
                            student=student,
                            academic_session=session_active,
                            semester_id=rec['sem_id']
                        )
                    
                    # Clean session cache
                    del request.session['validated_import_students']
                    
                    AuditLog.objects.create(
                        institution=inst,
                        user=request.user,
                        user_role=request.user.get_role_display(),
                        action=f"Bulk imported {len(records)} students via Excel/CSV.",
                        module="STUDENTS",
                        ip_address=request.META.get('REMOTE_ADDR')
                    )
                    
                messages.success(request, f"Successfully imported {len(records)} students!")
                return redirect('student_list')
                
            except Exception as e:
                messages.error(request, f"Import process failed: {e}")
                return redirect('student_import')
                
    return render(request, 'students/student_import.html', {
        'stage': stage,
        'errors': errors,
        'preview_records': preview_records,
        'breadcrumbs': [{'name': 'Students', 'url': '/students/'}, {'name': 'Import Students', 'url': '#'}]
    })
